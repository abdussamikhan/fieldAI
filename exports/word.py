import io
from typing import List, Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def _configure_inter_styles(doc: Document):
    """Enforces Inter font and non-bold styling across default Word styles."""
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Inter'
    font.size = Pt(10.5)
    font.bold = False

    for s_name in ['Heading 1', 'Heading 2', 'Heading 3', 'Title']:
        if s_name in doc.styles:
            h_style = doc.styles[s_name]
            h_style.font.name = 'Inter'
            h_style.font.bold = False

def export_transcript_word(
    title: str,
    summary_data: Dict[str, Any],
    segments: List[Dict[str, Any]]
) -> bytes:
    doc = Document()
    _configure_inter_styles(doc)
    
    # Title
    p_title = doc.add_paragraph()
    run_title = p_title.add_run(f"FieldAI Walkthrough Record: {title}")
    run_title.font.name = "Inter"
    run_title.font.size = Pt(18)
    run_title.font.bold = False
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    # Executive Summary Heading
    h1 = doc.add_heading("1. Executive Summary", level=1)
    for r in h1.runs:
        r.font.name = "Inter"
        r.font.bold = False
    p_sum = doc.add_paragraph(summary_data.get("summary", "No summary recorded."))
    for r in p_sum.runs:
        r.font.name = "Inter"
        r.font.bold = False

    # Key Points
    if summary_data.get("key_points"):
        h2 = doc.add_heading("2. Key Discussion Points", level=2)
        for r in h2.runs:
            r.font.name = "Inter"
            r.font.bold = False
        for pt in summary_data["key_points"]:
            p = doc.add_paragraph(pt, style='List Bullet')
            for r in p.runs:
                r.font.name = "Inter"
                r.font.bold = False

    # Open Questions
    if summary_data.get("open_questions"):
        h3 = doc.add_heading("3. Open Questions / Follow-ups", level=2)
        for r in h3.runs:
            r.font.name = "Inter"
            r.font.bold = False
        for q in summary_data["open_questions"]:
            p = doc.add_paragraph(q, style='List Bullet')
            for r in p.runs:
                r.font.name = "Inter"
                r.font.bold = False

    # PBC Requests
    if summary_data.get("pbc_requests"):
        h4 = doc.add_heading("4. Provided By Client (PBC) Requests", level=2)
        for r in h4.runs:
            r.font.name = "Inter"
            r.font.bold = False
        pbc_table = doc.add_table(rows=1, cols=3)
        pbc_table.style = 'Light Shading Accent 1'
        hdr_cells = pbc_table.rows[0].cells
        hdr_cells[0].text = "Requested Item"
        hdr_cells[1].text = "Owner"
        hdr_cells[2].text = "Due Date Hint"
        
        for cell in hdr_cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Inter"
                    r.font.bold = False

        for pbc in summary_data["pbc_requests"]:
            row_cells = pbc_table.add_row().cells
            row_cells[0].text = pbc.get("item", "")
            row_cells[1].text = pbc.get("owner", "")
            row_cells[2].text = pbc.get("due_hint", "")
            for cell in row_cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = "Inter"
                        r.font.bold = False

    # Detailed Transcript
    doc.add_page_break()
    h5 = doc.add_heading("5. Detailed Meeting Transcript", level=1)
    for r in h5.runs:
        r.font.name = "Inter"
        r.font.bold = False

    for seg in segments:
        speaker = seg.get("speaker_name") or seg.get("speaker_label", "Speaker")
        start_min = int(seg.get("start_s", 0) // 60)
        start_sec = int(seg.get("start_s", 0) % 60)
        time_str = f"[{start_min:02d}:{start_sec:02d}]"

        p = doc.add_paragraph()
        run_speaker = p.add_run(f"{speaker} {time_str}: ")
        run_speaker.font.name = "Inter"
        run_speaker.font.bold = False
        run_speaker.font.color.rgb = RGBColor(2, 132, 199)
        run_text = p.add_run(seg.get("text", ""))
        run_text.font.name = "Inter"
        run_text.font.bold = False

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()

def export_findings_word(findings: List[Dict[str, Any]]) -> bytes:
    doc = Document()
    _configure_inter_styles(doc)
    
    p_title = doc.add_paragraph()
    run_title = p_title.add_run("Internal Audit Fieldwork Findings (5 Cs)")
    run_title.font.name = "Inter"
    run_title.font.size = Pt(18)
    run_title.font.bold = False

    for idx, f in enumerate(findings, 1):
        h1 = doc.add_heading(f"Finding {idx}: {f.get('finding_code', '')} – {f.get('title', '')}", level=1)
        for r in h1.runs:
            r.font.name = "Inter"
            r.font.bold = False
        
        p_rating = doc.add_paragraph()
        run_r = p_rating.add_run(f"Severity Rating: {f.get('rating', 'Medium')}")
        run_r.font.name = "Inter"
        run_r.font.bold = False
        
        # 5 Cs
        for heading_label, text_key in [
            ("Condition (What was found):", "condition_text"),
            ("Criteria (What should be):", "criteria_text"),
            ("Cause (Why did it happen):", "cause_text"),
            ("Effect (Risk or business impact):", "effect_text"),
            ("Recommendation (Remediation action):", "recommendation_text")
        ]:
            h3 = doc.add_heading(heading_label, level=3)
            for r in h3.runs:
                r.font.name = "Inter"
                r.font.bold = False
            p_c = doc.add_paragraph(f.get(text_key) or f.get(text_key.replace("_text", ""), ""))
            for r in p_c.runs:
                r.font.name = "Inter"
                r.font.bold = False

        p_links = doc.add_paragraph(f"Linked Controls/Tests: {f.get('linked_controls', '')} / {f.get('linked_tests', '')}")
        for r in p_links.runs:
            r.font.name = "Inter"
            r.font.bold = False
        doc.add_paragraph("-" * 40)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
