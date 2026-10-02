import io
from typing import List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def _apply_header_style(ws, cols: List[str]):
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    for col_idx, col_name in enumerate(cols, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 28

def export_process_table_excel(steps: List[Dict[str, Any]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Process Steps"

    headers = [
        "Step Code", "Order", "Description", "Responsible Role", 
        "Department", "System", "Inputs", "Outputs", "Frequency", "Decision Point", "Confidence"
    ]
    _apply_header_style(ws, headers)

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for row_idx, step in enumerate(steps, 2):
        ws.cell(row=row_idx, column=1, value=step.get("step_code", ""))
        ws.cell(row=row_idx, column=2, value=step.get("order_num", 1))
        ws.cell(row=row_idx, column=3, value=step.get("description", ""))
        ws.cell(row=row_idx, column=4, value=step.get("responsible_role", ""))
        ws.cell(row=row_idx, column=5, value=step.get("department", ""))
        ws.cell(row=row_idx, column=6, value=step.get("system", ""))
        ws.cell(row=row_idx, column=7, value=step.get("inputs", ""))
        ws.cell(row=row_idx, column=8, value=step.get("outputs", ""))
        ws.cell(row=row_idx, column=9, value=step.get("frequency", "per transaction"))
        ws.cell(row=row_idx, column=10, value="Yes" if step.get("is_decision") else "No")
        ws.cell(row=row_idx, column=11, value=f"{int(step.get('confidence', 1.0) * 100)}%")

        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=c)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def export_rcm_excel(rcm_rows: List[Dict[str, Any]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Risk & Control Matrix"

    headers = [
        "Risk Code", "Risk Description", "Inherent Rating",
        "Control Code", "Control Description", "Control Type",
        "Nature", "Frequency", "Control Owner", "Key Control",
        "Design Rating", "Criteria / SOP Ref", "Framework Mapping"
    ]
    _apply_header_style(ws, headers)

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for row_idx, r in enumerate(rcm_rows, 2):
        ws.cell(row=row_idx, column=1, value=r.get("risk_code", ""))
        ws.cell(row=row_idx, column=2, value=r.get("risk_description", ""))
        ws.cell(row=row_idx, column=3, value=r.get("inherent_rating", "Medium"))
        ws.cell(row=row_idx, column=4, value=r.get("control_code", ""))
        ws.cell(row=row_idx, column=5, value=r.get("control_description", ""))
        ws.cell(row=row_idx, column=6, value=r.get("control_type", "preventive"))
        ws.cell(row=row_idx, column=7, value=r.get("nature", "automated"))
        ws.cell(row=row_idx, column=8, value=r.get("frequency", "per transaction"))
        ws.cell(row=row_idx, column=9, value=r.get("owner", ""))
        ws.cell(row=row_idx, column=10, value="Yes" if r.get("key_control", True) else "No")
        ws.cell(row=row_idx, column=11, value=r.get("design_rating", "Adequate"))
        ws.cell(row=row_idx, column=12, value=r.get("criteria_ref", ""))
        ws.cell(row=row_idx, column=13, value=", ".join(r.get("framework_refs", [])))

        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=c)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 45)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def export_audit_program_excel(tests: List[Dict[str, Any]]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Audit Program"

    headers = [
        "Test Code", "Control Code", "Test Objective", "Test Type",
        "Audit Procedure", "Sample Size", "Sample Basis", "Evidence / PBC",
        "Performer", "Reviewer", "WP Ref", "Result", "Exceptions / Notes"
    ]
    _apply_header_style(ws, headers)

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for row_idx, t in enumerate(tests, 2):
        ws.cell(row=row_idx, column=1, value=t.get("test_code", ""))
        ws.cell(row=row_idx, column=2, value=t.get("control_code", ""))
        ws.cell(row=row_idx, column=3, value=t.get("objective", ""))
        ws.cell(row=row_idx, column=4, value=t.get("test_type", "ToE"))
        ws.cell(row=row_idx, column=5, value=t.get("procedure", ""))
        ws.cell(row=row_idx, column=6, value=t.get("sample_size", 25))
        ws.cell(row=row_idx, column=7, value=t.get("sample_basis", "Frequency-based sampling table"))
        ws.cell(row=row_idx, column=8, value=t.get("evidence_pbc", ""))
        ws.cell(row=row_idx, column=9, value=t.get("performer", "Auditor"))
        ws.cell(row=row_idx, column=10, value=t.get("reviewer", "Audit Manager"))
        ws.cell(row=row_idx, column=11, value=t.get("wp_ref", ""))
        ws.cell(row=row_idx, column=12, value=t.get("result", "Untested"))
        ws.cell(row=row_idx, column=13, value=t.get("exception_note", ""))

        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=c)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 45)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
