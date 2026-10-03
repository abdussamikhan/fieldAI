"""
Generates a modern, widescreen (16:9) PowerPoint presentation for FieldAI hackathon pitch.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path="FieldAI_Hackathon_Pitch.pptx"):
    prs = Presentation()
    # Set 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]  # Blank

    # Color Palette
    BG_COLOR = RGBColor(15, 23, 42)       # Slate 900
    CARD_BG = RGBColor(30, 41, 59)        # Slate 800
    BORDER_COLOR = RGBColor(51, 65, 85)   # Slate 700
    TEXT_LIGHT = RGBColor(248, 250, 252)  # White / Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)  # Slate 400
    ACCENT_BLUE = RGBColor(56, 189, 248)  # Sky 400
    ACCENT_CYAN = RGBColor(6, 182, 212)   # Cyan 500

    def add_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="HACKATHON PROJECT PITCH"):
        # Category Tracker
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_CYAN
        p_cat.font.name = "Arial"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT
        p.font.name = "Arial"

    # ==========================================================
    # SLIDE 1: TITLE SLIDE
    # ==========================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_slide_background(s1)

    t_box = s1.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(11), Inches(3.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p0 = tf1.paragraphs[0]
    p0.text = "FIELD AI"
    p0.font.size = Pt(44)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_BLUE
    p0.font.name = "Arial"

    p1 = tf1.add_paragraph()
    p1.text = "Autonomous Multi-Agent Fieldwork & Process Intelligence"
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT
    p1.font.name = "Arial"
    p1.space_before = Pt(8)

    p2 = tf1.add_paragraph()
    p2.text = "Turning hours of messy walkthrough interviews & SOPs into verifiable audit workpapers in minutes."
    p2.font.size = Pt(14)
    p2.font.color.rgb = TEXT_MUTED
    p2.font.name = "Arial"
    p2.space_before = Pt(14)

    p3 = tf1.add_paragraph()
    p3.text = "Presented by: Sami Associates · Hackathon 2026"
    p3.font.size = Pt(12)
    p3.font.bold = True
    p3.font.color.rgb = ACCENT_CYAN
    p3.font.name = "Arial"
    p3.space_before = Pt(24)

    # ==========================================================
    # SLIDE 2: THE PROBLEM
    # ==========================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_background(s2)
    add_header(s2, "The Problem: The Fieldwork Bottleneck")

    # Card 1
    c1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.5))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = BORDER_COLOR
    tf_c1 = c1.text_frame
    tf_c1.word_wrap = True
    tf_c1.margin_left = Inches(0.3)
    tf_c1.margin_right = Inches(0.3)
    tf_c1.margin_top = Inches(0.3)

    p = tf_c1.paragraphs[0]
    p.text = "⏳  50%+ Audit Fieldwork Time Wasted"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    points_1 = [
        "Auditors spend days manually transcribing audio walkthrough interviews.",
        "Tedious re-keying across narratives, Visio flowcharts, and Risk & Control Matrices.",
        "Manual error and fatigue lead to incomplete risk evaluations."
    ]
    for pt in points_1:
        p = tf_c1.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(14)

    # Card 2
    c2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.5))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = BORDER_COLOR
    tf_c2 = c2.text_frame
    tf_c2.word_wrap = True
    tf_c2.margin_left = Inches(0.3)
    tf_c2.margin_right = Inches(0.3)
    tf_c2.margin_top = Inches(0.3)

    p = tf_c2.paragraphs[0]
    p.text = "🧩  The 'Say vs. Do' Gap & Broken Links"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    points_2 = [
        "What employees say in interviews often diverges from written SOP manuals.",
        "Identifying informal workarounds manually is exhausting and frequently missed.",
        "Traditional workpapers lack direct verifiable links back to exact audio timestamps or document clauses."
    ]
    for pt in points_2:
        p = tf_c2.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(14)

    # ==========================================================
    # SLIDE 3: THE SOLUTION
    # ==========================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_background(s3)
    add_header(s3, "The Solution: FieldAI Multi-Agent System")

    features = [
        ("🎙️ Walkthrough Ingestion", "Uploads multi-speaker audio; extracts accurate speech with speaker diarization & exact timestamps."),
        ("📄 Procedural Intelligence", "Ingests complex PDF & Word policies; automatically breaks them into hierarchical clauses."),
        ("🔍 Divergence & Gap Engine", "Reconciles interviews against SOPs across 5 divergence categories to detect operational drift."),
        ("⚡ 1-Click Workpaper Synthesis", "Generates standardized Process Tables, Mermaid Flowcharts, RCM, and Audit Programs with 100% citations.")
    ]

    for i, (title, desc) in enumerate(features):
        col = i % 2
        row = i // 2
        x = Inches(0.8 + col * 6.0)
        y = Inches(1.8 + row * 2.5)

        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.7), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12.5)
        p2.font.color.rgb = TEXT_LIGHT
        p2.space_before = Pt(8)

    # ==========================================================
    # SLIDE 4: ARCHITECTURE
    # ==========================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_background(s4)
    add_header(s4, "Multi-Agent Architecture: 21 LangGraph Agents")

    arch_points = [
        ("Master Orchestrator", "Central routing graph directing interview ingestion, SOP validation, and synthesis workflows."),
        ("Specialized Sub-Graphs", "6 coordinated sub-graphs: Transcription, Document QA, Process Flow, RCM, Audit Program, & Analytics."),
        ("Deterministic Schemas", "Every agent produces strictly typed JSON outputs validated before moving to downstream nodes."),
        ("Auditor in the Loop", "Supervisory review checkpoints allow auditors to inspect, edit, and approve before final compilation.")
    ]

    for i, (title, desc) in enumerate(arch_points):
        y = Inches(1.8 + i * 1.25)
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.7), Inches(1.05))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.18)

        p = tf.paragraphs[0]
        p.text = f"Agent Cluster: {title}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_LIGHT
        p2.space_before = Pt(3)

    # ==========================================================
    # SLIDE 5: TECH STACK
    # ==========================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_background(s5)
    add_header(s5, "Tech Stack & Engineering Highlights")

    techs = [
        ("LangGraph StateGraph", "Orchestration & State Machine", "Multi-agent acyclic graphs with checkpointing and state persistence."),
        ("DeepSeek Flash", "AI Reasoning Engine", "High-speed reasoning, clause alignment, and JSON schema compliance."),
        ("ElevenLabs Scribe", "Speech Intelligence", "Audio transcription with speaker diarization and millisecond-grade timestamps."),
        ("Streamlit Workstation", "Auditor UI Portal", "15-page interactive dashboard for reviewing evidence and testing matrices."),
        ("PostgreSQL & Jobs", "Data Persistence & Queues", "Background job worker handling asynchronous pipeline execution."),
        ("Audit Analytics Engine", "Deterministic Verification", "Full-population testing for duplicate payments, split POs, and SoD toxic roles.")
    ]

    for i, (name, role, detail) in enumerate(techs):
        col = i % 3
        row = i // 3
        x = Inches(0.8 + col * 4.0)
        y = Inches(1.8 + row * 2.5)

        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.7), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = name
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE

        p2 = tf.add_paragraph()
        p2.text = role
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = ACCENT_CYAN
        p2.space_before = Pt(2)

        p3 = tf.add_paragraph()
        p3.text = detail
        p3.font.size = Pt(11)
        p3.font.color.rgb = TEXT_LIGHT
        p3.space_before = Pt(6)

    # ==========================================================
    # SLIDE 6: IMPACT & DEMO
    # ==========================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_background(s6)
    add_header(s6, "Hackathon Results & Live Demo Takeaways")

    # Left Card
    c_res = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.5))
    c_res.fill.solid()
    c_res.fill.fore_color.rgb = CARD_BG
    c_res.line.color.rgb = BORDER_COLOR
    tf = c_res.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.3)

    p = tf.paragraphs[0]
    p.text = "🏆 Measurable Impact"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    res_pts = [
        ("> 50% Reduction in Fieldwork Hours", "Saves 3-5 days per audit cycle on administrative drafting."),
        ("100% Traceability & Compliance", "Every finding and test directly cites source audio or policy clauses."),
        ("Full-Population Testing", "Replaces standard 25-item samples with 100% automated substantive checks.")
    ]
    for bold_text, sub_text in res_pts:
        p = tf.add_paragraph()
        p.text = f"• {bold_text}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(12)

        p2 = tf.add_paragraph()
        p2.text = f"  {sub_text}"
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(2)

    # Right Card
    c_road = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.5))
    c_road.fill.solid()
    c_road.fill.fore_color.rgb = CARD_BG
    c_road.line.color.rgb = BORDER_COLOR
    tf2 = c_road.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.3)
    tf2.margin_top = Inches(0.3)

    p = tf2.paragraphs[0]
    p.text = "🚀 Future Vision"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    road_pts = [
        ("Offline-First Fieldwork Sync", "Enable full local processing on encrypted laptops in remote client sites."),
        ("Real-time Interview Co-Pilot", "Live suggestions for auditors to ask dynamic probing follow-up questions."),
        ("Direct ERP Ingestion", "Seamless API connectors for SAP S/4HANA, NetSuite, and Workday event logs.")
    ]
    for bold_text, sub_text in road_pts:
        p = tf2.add_paragraph()
        p.text = f"• {bold_text}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(12)

        p2 = tf2.add_paragraph()
        p2.text = f"  {sub_text}"
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(2)

    # ==========================================================
    # SLIDE 7: CONCLUSION
    # ==========================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_background(s7)

    t_box = s7.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.3), Inches(3.0))
    tf7 = t_box.text_frame
    tf7.word_wrap = True
    
    p = tf7.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "Thank You! 🎯"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.font.name = "Arial"

    p2 = tf7.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "FieldAI: Reimagining Audit Fieldwork with Multi-Agent AI"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_LIGHT
    p2.space_before = Pt(12)

    p3 = tf7.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    p3.text = "Ready for Q&A and Live Demonstration"
    p3.font.size = Pt(14)
    p3.font.color.rgb = ACCENT_CYAN
    p3.space_before = Pt(18)

    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck("FieldAI_Hackathon_Pitch.pptx")
